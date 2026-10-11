"""Parser Bank Soal USKP (.docx) -> PG + esai, beserta validasi untuk pratinjau impor."""
import re
import uuid
from collections import Counter

import docx

ANSWER_MARKERS = ("Kunci Jawaban", "Pedoman Jawaban", "Jawaban:")
STOP_MARKERS = ("Catatan Jebakan", "Tingkat Kesulitan", "Tingkat Perhitungan", "Verifikasi")
KEYS = ["A", "B", "C", "D"]


def _doc(src):
    return src if hasattr(src, "paragraphs") else docx.Document(src)


def short_materi(text):
    m = re.match(r"^(.*?)\s*\(([^()]*)\)\s*$", text)
    if m and re.search(r"\d|UU|PMK|Pasal|PER|PP|SE-", m.group(2)):
        return m.group(1).strip()
    return text.strip()


def parse(src):
    d = _doc(src)
    paras = [p.text.strip() for p in d.paragraphs]

    mcq, essays = [], []
    mode = "mcq"
    cur = None
    phase = field = last_opt = None
    materi = ""

    def push_mcq():
        nonlocal cur
        if cur and mode == "mcq":
            mcq.append(cur)
        cur = None

    def push_essay():
        nonlocal cur
        if cur and mode == "essay":
            essays.append(cur)
        cur = None

    for t in paras:
        if not t:
            continue
        if t.startswith("Bagian III"):
            push_essay(); break
        if t.startswith("Bagian II"):
            push_mcq(); mode = "essay"; cur = None; continue

        if mode == "mcq" and re.match(r"^Esai\s+\d+\s*[\u2014\u2013\-]", t):
            push_mcq(); mode = "essay"; cur = None

        if mode == "mcq":
            if t.startswith("Materi"):
                push_mcq()
                materi = short_materi(t.split(":", 1)[1]) if ":" in t else ""
                field = None; continue
            m = re.match(r"^Soal\s+(\d+)\.\s*(.*)$", t)
            if m:
                push_mcq()
                cur = {"number": int(m.group(1)), "question": m.group(2).strip(), "materi": materi,
                       "options": {}, "answer": None, "pembahasan": "", "dasar": ""}
                phase = "stem"; field = None; last_opt = None
                continue
            if not cur:
                continue
            mo = re.match(r"^([ABCD])\.\s+(.*)$", t)
            if mo:
                cur["options"][mo.group(1)] = mo.group(2).strip()
                last_opt = mo.group(1); phase = "options"; field = None
                continue
            if t.startswith("Kunci Jawaban"):
                mk = re.search(r"([ABCD])", t.split(":", 1)[1] if ":" in t else t)
                cur["answer"] = mk.group(1) if mk else None
                phase = "post"; field = None
                continue
            if t.startswith("Pembahasan"):
                cur["pembahasan"] = t.split(":", 1)[1].strip() if ":" in t else ""
                field = "pembahasan"; continue
            if t.startswith("Dasar Hukum"):
                cur["dasar"] = t.split(":", 1)[1].strip() if ":" in t else ""
                field = "dasar"; continue
            if t.startswith(STOP_MARKERS):
                field = None; continue
            if phase == "stem":
                cur["question"] += " " + t
            elif phase == "options" and last_opt:
                cur["options"][last_opt] += " " + t
            elif field == "pembahasan":
                cur["pembahasan"] += " " + t
            elif field == "dasar":
                cur["dasar"] += " " + t
        else:
            m = re.match(r"^Esai\s+(\d+)\s*[\u2014\u2013\-]\s*(.*)$", t)
            if m:
                push_essay()
                cur = {"number": int(m.group(1)), "topic": m.group(2).strip(),
                       "question": "", "model": "", "dasar": ""}
                field = "question"; continue
            if not cur:
                continue
            if t.startswith("Kasus"):
                rest = t.split(":", 1)[1].strip() if ":" in t else ""
                cur["question"] = (cur["question"] + "\n" + rest).strip() if cur["question"] else rest
                field = "question"; continue
            if t.startswith(ANSWER_MARKERS):
                cur["model"] = t.split(":", 1)[1].strip() if ":" in t else ""
                field = "model"; continue
            if t.startswith("Dasar Hukum"):
                cur["dasar"] = t.split(":", 1)[1].strip() if ":" in t else ""
                field = "dasar"; continue
            if t.startswith(STOP_MARKERS):
                field = None; continue
            if field == "question":
                cur["question"] += ("\n" if cur["question"] else "") + t
            elif field == "model":
                cur["model"] += ("\n" if cur["model"] else "") + t
            elif field == "dasar":
                cur["dasar"] += " " + t

    push_mcq(); push_essay()
    return mcq, essays


def read_meta(d):
    level, name = "", ""
    for p in d.paragraphs[:15]:
        t = p.text.strip()
        ml = re.search(r"TINGKAT\s+([ABC])\b", t, re.I)
        if ml and not level:
            level = ml.group(1).upper()
        if t.lower().startswith("mata ujian") and ":" in t and not name:
            name = t.split(":", 1)[1].strip()
    return {"level": level, "name": name}


def key_table(d):
    for t in d.tables:
        head = [c.text.strip().lower() for c in t.rows[0].cells]
        if "no" in head and "kunci" in head:
            key = {}
            for r in t.rows[1:]:
                c = [x.text.strip() for x in r.cells]
                for i in range(0, len(c) - 1, 2):
                    if c[i].isdigit() and c[i + 1] in KEYS:
                        key[int(c[i])] = c[i + 1]
            return key
    return {}


def analyze(src):
    d = _doc(src)
    mcq, essays = parse(d)
    keys = key_table(d)
    issues, warnings, valid_mcq, valid_essays = [], [], [], []

    seen = set()
    for q in mcq:
        n = q["number"]
        missing = [k for k in KEYS if not (q["options"].get(k) or "").strip()]
        if n in seen:
            issues.append({"type": "PG", "number": n, "message": "Nomor soal ganda (duplikat)"})
            continue
        seen.add(n)
        if not q["question"].strip():
            issues.append({"type": "PG", "number": n, "message": "Teks pertanyaan kosong"})
        elif missing:
            issues.append({"type": "PG", "number": n, "message": f"Opsi {', '.join(missing)} tidak ditemukan"})
        elif q["answer"] not in KEYS:
            issues.append({"type": "PG", "number": n, "message": "Kunci jawaban tidak ditemukan"})
        else:
            valid_mcq.append(q)
            if keys and keys.get(n) and keys[n] != q["answer"]:
                warnings.append({"type": "PG", "number": n,
                                 "message": f"Kunci di soal = {q['answer']}, di tabel ringkasan = {keys[n]} (dipakai: {q['answer']})"})
    if seen:
        gaps = sorted(set(range(1, max(seen) + 1)) - seen)
        for n in gaps:
            warnings.append({"type": "PG", "number": n, "message": "Nomor soal tidak ditemukan di dokumen"})
    no_materi = sum(1 for q in valid_mcq if not q["materi"])
    if valid_mcq and no_materi:
        warnings.append({"type": "PG", "number": None, "message": f"{no_materi} soal tanpa judul 'Materi:' (tidak masuk rapor per materi)"})

    eseen = set()
    for e in essays:
        n = e["number"]
        if n in eseen:
            issues.append({"type": "Esai", "number": n, "message": "Nomor esai ganda (duplikat)"})
            continue
        eseen.add(n)
        if not e["question"].strip():
            issues.append({"type": "Esai", "number": n, "message": "Kasus/pertanyaan esai kosong"})
        elif not e["model"].strip():
            issues.append({"type": "Esai", "number": n, "message": "Jawaban model esai kosong"})
        else:
            valid_essays.append(e)

    materi = Counter(q["materi"] for q in valid_mcq if q["materi"])
    return {
        "meta": read_meta(d),
        "mcq_total": len(mcq),
        "essay_total": len(essays),
        "valid_mcq": valid_mcq,
        "valid_essays": valid_essays,
        "issues": issues,
        "warnings": warnings,
        "materi": [{"name": k, "count": v} for k, v in materi.items()],
        "key_table_found": bool(keys),
    }


def to_qdocs(sid, mcq):
    out = []
    for q in mcq:
        pemb = (q.get("pembahasan") or "").strip()
        dasar = (q.get("dasar") or "").strip()
        if dasar:
            pemb = (pemb + "\n\nDasar Hukum: " + dasar).strip()
        out.append({
            "id": str(uuid.uuid4()), "subject_id": sid, "number": q["number"],
            "question": q["question"].strip(), "materi": q.get("materi", ""),
            "options": [{"key": k, "text": q["options"][k].strip()} for k in KEYS],
            "answer": q["answer"], "pembahasan": pemb,
        })
    return out


def to_edocs(sid, essays):
    out = []
    for e in essays:
        model = (e.get("model") or "").strip()
        dasar = (e.get("dasar") or "").strip()
        if dasar:
            model = (model + "\n\nDasar Hukum: " + dasar).strip()
        topic = (e.get("topic") or "").strip()
        out.append({
            "id": str(uuid.uuid4()), "subject_id": sid, "number": e["number"],
            "question": (f"[{topic}] " if topic else "") + e["question"].strip(),
            "model_answer": model,
        })
    return out
